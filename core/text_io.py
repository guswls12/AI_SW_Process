"""소스 파일 텍스트 읽기 — 한국 윈도우 환경 인코딩 폴백 처리.

- UTF-8(BOM 포함) → CP949 → latin-1 순으로 시도.
- 모든 인코딩에서 latin-1은 실패하지 않으므로 최종 폴백.
"""

import os

_ENCODINGS = ("utf-8-sig", "cp949", "latin-1")


def read_text(path: str) -> str:
    """파일을 텍스트로 읽어 문자열 반환. 인코딩 자동 폴백."""
    with open(path, "rb") as f:
        data = f.read()
    for enc in _ENCODINGS:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def read_text_lines(path: str) -> list:
    """파일을 텍스트로 읽어 줄 단위 리스트 반환 (줄 끝 \\r\\n 제거)."""
    return [l.rstrip("\r\n") for l in read_text(path).splitlines()]


def read_doc_lines(path: str) -> list:
    """요구사항 문서 (.docx/.md/.txt) 를 줄 단위 리스트로 읽는다.

    .docx 는 paragraphs + tables 를 순서대로 평탄화 (빈 줄 제외).
    그 외 확장자는 read_text_lines 로 폴백 (인코딩 자동).
    docx 패키지는 lazy import — 미설치 환경에선 ImportError 가 호출 측으로 전파된다.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        lines: list = []
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                lines.append(t)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        lines.append(t)
        return lines
    return read_text_lines(path)
