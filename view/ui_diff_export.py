"""ui_diff_export.py — 코드 변경점(DIFF) 내보내기 헬퍼 (Excel/HTML).

main.py 의 MainWindow 메서드(_write_xlsx_diff_sheet, _safe_xlsx_sheet_name,
_on_export_full_xlsx, _build_html_diff_rows, _build_html_file_section,
_on_export_full_html, _html_shell)를 모듈 함수로 분리한다.

MainWindow 는 본 모듈을 import 해서 export_full_xlsx() / export_full_html() 만
호출한다 — 부모 위젯과 result panel/status bar/old·new 라인을 인자로 넘긴다.

ui_save.py 가 'AI 분석 결과(MD/HTML/DOCX)' 저장 책임을 갖는 것과 분리되어,
본 모듈은 'DIFF 코드 변경점(Excel/HTML)' 내보내기만 담당한다.

성능 최적화:
- Excel 은 openpyxl write_only 모드 + NamedStyle 로 셀 단위 스타일 비용 제거
- 폴더 모드 다중 파일 저장은 ThreadPoolExecutor 로 병렬화
- QProgressDialog 로 진행률 표시 (UI 무응답 방지)
"""

import os
import re
import datetime
import difflib
from concurrent.futures import ThreadPoolExecutor, as_completed

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment, NamedStyle

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QFileDialog, QProgressDialog, QApplication


# ──────────────────────────────────────────────────────────────
#  NamedStyle 빌더 — 워크북마다 새로 만들어 등록 (NamedStyle 은 단일 워크북 귀속)
# ──────────────────────────────────────────────────────────────
def _build_diff_styles() -> dict:
    """diff 4열 표에 사용할 NamedStyle dict 반환. 색상/폰트/정렬은 기존과 동일."""
    eq_font    = Font(color="8EB4D6", size=10)
    eq_fill    = PatternFill("solid", fgColor="0D1B2A")
    add_font   = Font(color="34D399", size=10)
    add_fill   = PatternFill("solid", fgColor="082318")
    del_font   = Font(color="F87171", size=10)
    del_fill   = PatternFill("solid", fgColor="230808")
    mod_font_l = Font(color="FCD34D", size=10)
    mod_fill_l = PatternFill("solid", fgColor="1C1200")
    mod_font_r = Font(color="34D399", size=10)
    mod_fill_r = PatternFill("solid", fgColor="082318")
    center_align = Alignment(horizontal="center")

    hdr = NamedStyle(name="diff_hdr")
    hdr.font = Font(bold=True, color="FFFFFF", size=10)
    hdr.fill = PatternFill("solid", fgColor="1E3A5F")
    hdr.alignment = Alignment(horizontal="center", vertical="center")

    def _mk(name: str, font: Font, fill: PatternFill, center: bool = False) -> NamedStyle:
        ns = NamedStyle(name=name)
        ns.font = font
        ns.fill = fill
        if center:
            ns.alignment = center_align
        return ns

    return {
        "hdr":        hdr,
        "eq_ln":      _mk("diff_eq_ln",     eq_font,    eq_fill,    center=True),
        "eq_code":    _mk("diff_eq_code",   eq_font,    eq_fill),
        "add_ln":     _mk("diff_add_ln",    add_font,   add_fill,   center=True),
        "add_code":   _mk("diff_add_code",  add_font,   add_fill),
        "del_ln":     _mk("diff_del_ln",    del_font,   del_fill,   center=True),
        "del_code":   _mk("diff_del_code",  del_font,   del_fill),
        "mod_l_ln":   _mk("diff_mod_l_ln",  mod_font_l, mod_fill_l, center=True),
        "mod_l_code": _mk("diff_mod_l_code",mod_font_l, mod_fill_l),
        "mod_r_ln":   _mk("diff_mod_r_ln",  mod_font_r, mod_fill_r, center=True),
        "mod_r_code": _mk("diff_mod_r_code",mod_font_r, mod_fill_r),
    }


def _register_styles(wb) -> None:
    """워크북에 diff 용 NamedStyle 들을 한 번에 등록."""
    for ns in _build_diff_styles().values():
        wb.add_named_style(ns)


# ──────────────────────────────────────────────────────────────
#  Excel diff 시트 헬퍼 (write_only + NamedStyle)
# ──────────────────────────────────────────────────────────────
def write_xlsx_diff_sheet(ws, old_lines, new_lines):
    """write_only 워크시트에 old/new 의 diff 를 4열로 기록.
    호출 측은 사전에 wb 에 _register_styles(wb) 로 NamedStyle 들을 등록해야 함.
    """
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 55

    def _c(value, style_name: str):
        c = WriteOnlyCell(ws, value=value)
        c.style = style_name
        return c

    ws.append([
        _c("구 행#",      "diff_hdr"),
        _c("수정 전 코드", "diff_hdr"),
        _c("신 행#",      "diff_hdr"),
        _c("수정 후 코드", "diff_hdr"),
    ])
    ws.row_dimensions[1].height = 20

    matcher = difflib.SequenceMatcher(None, old_lines, new_lines, autojunk=False)
    on = nn = 0
    append = ws.append  # 지역 바인딩으로 미세 가속

    for op, a0, a1, b0, b1 in matcher.get_opcodes():
        if op == "equal":
            for i in range(a1 - a0):
                on += 1; nn += 1
                append([
                    _c(on,                "diff_eq_ln"),
                    _c(old_lines[a0 + i], "diff_eq_code"),
                    _c(nn,                "diff_eq_ln"),
                    _c(new_lines[b0 + i], "diff_eq_code"),
                ])
        elif op == "replace":
            oc = old_lines[a0:a1]; nc = new_lines[b0:b1]
            for i in range(max(len(oc), len(nc))):
                ol = oc[i] if i < len(oc) else None
                nl = nc[i] if i < len(nc) else None
                if ol is not None: on += 1
                if nl is not None: nn += 1
                append([
                    _c(on if ol is not None else "", "diff_mod_l_ln"),
                    _c(ol or "",                     "diff_mod_l_code"),
                    _c(nn if nl is not None else "", "diff_mod_r_ln"),
                    _c(nl or "",                     "diff_mod_r_code"),
                ])
        elif op == "delete":
            for i in range(a1 - a0):
                on += 1
                append([
                    _c(on,                "diff_del_ln"),
                    _c(old_lines[a0 + i], "diff_del_code"),
                    _c("",                "diff_eq_ln"),
                    _c("",                "diff_eq_code"),
                ])
        elif op == "insert":
            for i in range(b1 - b0):
                nn += 1
                append([
                    _c("",                "diff_eq_ln"),
                    _c("",                "diff_eq_code"),
                    _c(nn,                "diff_add_ln"),
                    _c(new_lines[b0 + i], "diff_add_code"),
                ])


def safe_xlsx_sheet_name(base: str, used: set) -> str:
    """엑셀 시트명 제약(31자 이하, \\/ ?*[]: 제외, 중복 불가)에 맞춰 정규화."""
    cleaned = re.sub(r'[\\/?*\[\]:]', '_', base)[:31] or "Sheet"
    name = cleaned; n = 2
    while name in used:
        suffix = f"_{n}"
        name = cleaned[:31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


# ──────────────────────────────────────────────────────────────
#  단일 파일 저장 (스레드에서 호출 가능 — Qt 의존 없음)
# ──────────────────────────────────────────────────────────────
def _save_one_diff_xlsx(rel_path: str, old_lines, new_lines, out_path: str) -> None:
    """단일 파일 diff 를 .xlsx 로 저장 (write_only)."""
    wb = openpyxl.Workbook(write_only=True)
    _register_styles(wb)
    title = safe_xlsx_sheet_name(os.path.basename(rel_path) or rel_path, set())
    ws = wb.create_sheet(title=title)
    write_xlsx_diff_sheet(ws, old_lines, new_lines)
    wb.save(out_path)


def _save_one_diff_html(rel_path: str, old_lines, new_lines, out_path: str, ts: str) -> None:
    """단일 파일 diff 를 .html 로 저장."""
    fname = rel_path.replace("\\", "/").split("/")[-1] or rel_path
    page_title = f"🔀 코드 변경점 — {fname}"
    rows_html_joined = build_html_diff_rows(old_lines, new_lines)
    body_block = f"""
<div class="wrap">
<table>
<colgroup><col><col><col><col></colgroup>
<thead>
  <tr>
    <th>행#</th><th>수정 전 코드</th>
    <th>행#</th><th>수정 후 코드</th>
  </tr>
</thead>
<tbody>
{rows_html_joined}
</tbody>
</table>
</div>"""
    html = html_shell(page_title, ts, body_block)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)


# ──────────────────────────────────────────────────────────────
#  병렬 저장 + 진행률 다이얼로그
# ──────────────────────────────────────────────────────────────
def _run_parallel_with_progress(parent, items, save_fn, kind_label: str) -> tuple:
    """items 각각에 save_fn(*item) 을 ThreadPoolExecutor 로 병렬 실행.
    QProgressDialog 로 진행률 표시 + 취소 처리.
    반환: (저장된 개수, 취소 여부, 실패 목록[(rel_path, str(error))])
    """
    n = len(items)
    dlg = QProgressDialog(f"{kind_label} 저장 중... (0/{n})", "취소", 0, n, parent)
    dlg.setWindowModality(Qt.WindowModality.WindowModal)
    dlg.setMinimumDuration(0)
    dlg.setValue(0)
    QApplication.processEvents()

    saved = 0
    failures: list = []
    cancelled = False
    max_workers = min(8, max(2, (os.cpu_count() or 2)))

    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        futures = {ex.submit(save_fn, *item): item for item in items}
        for fut in as_completed(futures):
            item = futures[fut]
            rel_path = item[0] if item else "?"
            if dlg.wasCanceled():
                cancelled = True
                for f in futures:
                    f.cancel()
                break
            try:
                fut.result()
                saved += 1
            except Exception as e:
                failures.append((rel_path, str(e)))
            done = saved + len(failures)
            dlg.setValue(done)
            dlg.setLabelText(f"{kind_label} 저장 중... ({done}/{n})")
            QApplication.processEvents()

    dlg.close()
    return saved, cancelled, failures


# ──────────────────────────────────────────────────────────────
#  HTML diff 헬퍼
# ──────────────────────────────────────────────────────────────
def build_html_diff_rows(old_lines, new_lines) -> str:
    """old/new 를 HTML tbody 의 <tr>...</tr> 문자열 연결로 반환."""
    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    rows_html = []
    matcher = difflib.SequenceMatcher(None, old_lines, new_lines, autojunk=False)
    on = nn = 0
    for op, a0, a1, b0, b1 in matcher.get_opcodes():
        if op == "equal":
            for i in range(a1 - a0):
                on += 1; nn += 1
                rows_html.append(
                    f'<tr class="eq">'
                    f'<td class="ln">{on}</td><td class="code">{esc(old_lines[a0+i])}</td>'
                    f'<td class="ln">{nn}</td><td class="code">{esc(new_lines[b0+i])}</td>'
                    f'</tr>')
        elif op == "replace":
            oc = old_lines[a0:a1]; nc = new_lines[b0:b1]
            for i in range(max(len(oc), len(nc))):
                ol = oc[i] if i < len(oc) else None
                nl = nc[i] if i < len(nc) else None
                if ol is not None: on += 1
                if nl is not None: nn += 1
                ln_l  = str(on) if ol is not None else ""
                ln_r  = str(nn) if nl is not None else ""
                rows_html.append(
                    f'<tr>'
                    f'<td class="ln mod_l">{ln_l}</td>'
                    f'<td class="code mod_l">{esc(ol or "")}</td>'
                    f'<td class="ln mod_r">{ln_r}</td>'
                    f'<td class="code mod_r">{esc(nl or "")}</td>'
                    f'</tr>')
        elif op == "delete":
            for i in range(a1 - a0):
                on += 1
                rows_html.append(
                    f'<tr>'
                    f'<td class="ln del">{on}</td>'
                    f'<td class="code del">{esc(old_lines[a0+i])}</td>'
                    f'<td class="ln eq"></td><td class="code eq"></td>'
                    f'</tr>')
        elif op == "insert":
            for i in range(b1 - b0):
                nn += 1
                rows_html.append(
                    f'<tr>'
                    f'<td class="ln eq"></td><td class="code eq"></td>'
                    f'<td class="ln add">{nn}</td>'
                    f'<td class="code add">{esc(new_lines[b0+i])}</td>'
                    f'</tr>')
    return ''.join(rows_html)


def build_html_file_section(rel_path: str, old_lines, new_lines) -> str:
    """파일 하나에 대한 HTML 섹션(제목 + table) 을 반환."""
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    rows = build_html_diff_rows(old_lines, new_lines)
    return f"""
<section class="file-block">
  <h2 class="file-title">📄 {esc(rel_path)}</h2>
  <div class="wrap">
  <table>
  <colgroup><col><col><col><col></colgroup>
  <thead>
    <tr>
      <th>행#</th><th>수정 전 코드</th>
      <th>행#</th><th>수정 후 코드</th>
    </tr>
  </thead>
  <tbody>
  {rows}
  </tbody>
  </table>
  </div>
</section>"""


def html_shell(page_title: str, ts: str, body_inner: str) -> str:
    """HTML 페이지의 공통 스타일과 헤더를 감싸는 셸.
    body_inner 에는 하나 이상의 diff 표 또는 섹션/TOC 가 들어간다."""
    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<title>{page_title} — {ts}</title>
<style>
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ background: #0A1628; color: #8EB4D6; font-family: 'Consolas','D2Coding',monospace; font-size: 13px; }}
  h1 {{ padding: 18px 24px 6px; font-size: 16px; color: #5B9BD5; border-bottom: 1px solid #1E3A5F; }}
  h2.file-title {{ padding: 14px 24px 6px; font-size: 14px; color: #5B9BD5;
                   border-top: 1px solid #1E3A5F; background: #0D1B2A; }}
  .meta {{ padding: 6px 24px 14px; font-size: 11px; color: #3D5570; }}
  .legend {{ display: flex; gap: 20px; padding: 10px 24px; border-bottom: 1px solid #1E3A5F; font-size: 11px; }}
  .legend span {{ display: flex; align-items: center; gap: 6px; }}
  .dot {{ width: 12px; height: 12px; border-radius: 2px; }}
  .wrap {{ overflow-x: auto; }}
  table {{ border-collapse: collapse; width: 100%; table-layout: fixed; }}
  colgroup col:nth-child(1), colgroup col:nth-child(3) {{ width: 52px; }}
  colgroup col:nth-child(2), colgroup col:nth-child(4) {{ width: calc(50% - 52px); }}
  thead th {{ background: #1E3A5F; color: #fff; padding: 6px 8px; font-size: 12px; text-align: center; position: sticky; top: 0; }}
  thead th:nth-child(2), thead th:nth-child(4) {{ text-align: left; padding-left: 12px; }}
  td {{ padding: 1px 4px; white-space: pre; overflow: hidden; text-overflow: ellipsis; vertical-align: middle; }}
  td.ln {{ text-align: right; color: #2A4A6A; font-size: 11px; padding-right: 6px; user-select: none; border-right: 1px solid #1E3A5F; }}
  td.code {{ padding-left: 10px; }}
  tr.eq td {{ background: #0D1B2A; color: #8EB4D6; }}
  tr.eq td.ln {{ color: #1E3A5F; }}
  td.add {{ background: #082318; color: #34D399; }}
  td.add.ln {{ color: #1A5C3A; }}
  td.del {{ background: #230808; color: #F87171; }}
  td.del.ln {{ color: #5C1A1A; }}
  td.mod_l {{ background: #1C1200; color: #FCD34D; }}
  td.mod_l.ln {{ color: #4A3300; }}
  td.mod_r {{ background: #082318; color: #34D399; }}
  td.mod_r.ln {{ color: #1A5C3A; }}
  tr:hover td {{ filter: brightness(1.15); }}
  nav.toc {{ padding: 12px 24px; background: #0D1B2A; border-bottom: 1px solid #1E3A5F; }}
  nav.toc h3 {{ color: #5B9BD5; font-size: 12px; margin-bottom: 8px; }}
  nav.toc ul {{ list-style: none; padding-left: 8px; columns: 2; }}
  nav.toc li {{ margin: 4px 0; font-size: 12px; }}
  nav.toc a {{ color: #8EB4D6; text-decoration: none; }}
  nav.toc a:hover {{ color: #5B9BD5; text-decoration: underline; }}
  section.file-block {{ margin-bottom: 18px; }}
</style>
</head>
<body>
<h1>{page_title}</h1>
<div class="meta">생성 일시: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</div>
<div class="legend">
  <span><span class="dot" style="background:#8EB4D6"></span>변경 없음</span>
  <span><span class="dot" style="background:#F87171"></span>삭제</span>
  <span><span class="dot" style="background:#FCD34D"></span>수정 전</span>
  <span><span class="dot" style="background:#34D399"></span>수정 후</span>
</div>
{body_inner}
</body>
</html>"""


# ──────────────────────────────────────────────────────────────
#  내보내기 진입점
# ──────────────────────────────────────────────────────────────
def _make_safe_fname_factory(ts: str):
    """폴더 모드에서 출력 파일명 충돌 방지 팩토리. (메인 스레드에서 사용)"""
    used_names: set = set()
    def _safe_fname(rel_path: str, ext: str) -> str:
        base = os.path.basename(rel_path) or rel_path
        stem, _src_ext = os.path.splitext(base)
        cleaned = re.sub(r'[\\/?*\[\]:"<>|]', '_', stem) or "diff"
        name = f"Diff_{cleaned}_{ts}{ext}"
        n = 2
        while name in used_names:
            name = f"Diff_{cleaned}_{ts}_{n}{ext}"
            n += 1
        used_names.add(name)
        return name
    return _safe_fname


def export_full_xlsx(parent, result_panel, status_bar, old_lines, new_lines):
    """전체 코드 변경점을 .xlsx 로 저장.

    - 폴더 모드 + 체크된 파일이 있으면: 체크된 파일 각각을 별도 .xlsx 로 (병렬) 저장
    - 그 외: 현재 선택된 단일 파일 또는 전체 병합 diff 를 한 파일로 저장
    """
    if not old_lines and not new_lines:
        status_bar.showMessage("⚠  먼저 DIFF 추출을 실행해주세요."); return

    # ── 폴더 모드 + 체크된 파일 있음 → 체크된 파일을 각각 별도 파일로 저장 ──
    checked_pairs = []
    if getattr(result_panel, "_folder_mode", False):
        checked_pairs = result_panel.get_checked_folder_pairs()

    if checked_pairs:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        out_dir = QFileDialog.getExistingDirectory(
            parent, f"체크된 {len(checked_pairs)}개 파일 저장 폴더 선택")
        if not out_dir:
            return

        safe_fname = _make_safe_fname_factory(ts)
        items = []
        for rel_path, old_lines_pair, new_lines_pair in checked_pairs:
            out_path = os.path.join(out_dir, safe_fname(rel_path, ".xlsx"))
            items.append((rel_path, old_lines_pair, new_lines_pair, out_path))

        saved, cancelled, failures = _run_parallel_with_progress(
            parent, items, _save_one_diff_xlsx, "엑셀")

        if cancelled:
            status_bar.showMessage(
                f"⚠  엑셀 저장 취소됨 ({saved}/{len(checked_pairs)}개 저장): {out_dir}")
        elif failures:
            status_bar.showMessage(
                f"⚠  엑셀 저장 일부 실패 ({saved}/{len(checked_pairs)}개 성공): {out_dir}")
        else:
            status_bar.showMessage(
                f"✅  엑셀 저장 완료 ({saved}개 파일, 개별 저장): {out_dir}")
        return

    # ── 기존 동작: 현재 선택된 단일 파일 or 전체 병합 ──
    cur = result_panel.get_current_file_lines()
    if cur:
        rel_path, old_lines, new_lines = cur
        fname = os.path.basename(rel_path) if hasattr(os, 'path') else rel_path.split("/")[-1].split("\\")[-1]
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Diff_{fname}_{ts}.xlsx"
        title_label = f"{fname} 변경점"
    else:
        rel_path = "전체 변경점"
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Full_Diff_{ts}.xlsx"
        title_label = "전체 변경점"

    path, _ = QFileDialog.getSaveFileName(
        parent, "코드 변경점 엑셀 저장", default_name, "Excel (*.xlsx)")
    if not path:
        return

    # 단일 파일도 진행률 표시 + 백그라운드 처리는 생략 (단일 파일은 빠름 — 메시지로만 안내)
    dlg = QProgressDialog("엑셀 저장 중...", None, 0, 0, parent)
    dlg.setWindowModality(Qt.WindowModality.WindowModal)
    dlg.setCancelButton(None)
    dlg.setMinimumDuration(0)
    dlg.show()
    QApplication.processEvents()
    try:
        wb = openpyxl.Workbook(write_only=True)
        _register_styles(wb)
        ws = wb.create_sheet(title=safe_xlsx_sheet_name(title_label, set()))
        write_xlsx_diff_sheet(ws, old_lines, new_lines)
        wb.save(path)
    finally:
        dlg.close()
    status_bar.showMessage(f"✅  엑셀 저장 완료: {path}")


def export_full_html(parent, result_panel, status_bar, old_lines, new_lines):
    """전체 코드 변경점을 .html 로 저장.

    - 폴더 모드 + 체크된 파일이 있으면: 체크된 파일 각각을 별도 .html 로 (병렬) 저장
    - 그 외: 현재 선택된 단일 파일 또는 전체 병합 diff 를 한 파일로 저장
    """
    if not old_lines and not new_lines:
        status_bar.showMessage("⚠  먼저 DIFF 추출을 실행해주세요."); return

    # ── 폴더 모드 + 체크된 파일 있음 → 파일별 별도 .html 로 저장 ──
    checked_pairs = []
    if getattr(result_panel, "_folder_mode", False):
        checked_pairs = result_panel.get_checked_folder_pairs()

    if checked_pairs:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        out_dir = QFileDialog.getExistingDirectory(
            parent, f"체크된 {len(checked_pairs)}개 파일 저장 폴더 선택")
        if not out_dir:
            return

        safe_fname = _make_safe_fname_factory(ts)
        items = []
        for rel_path, old_lines_pair, new_lines_pair in checked_pairs:
            out_path = os.path.join(out_dir, safe_fname(rel_path, ".html"))
            items.append((rel_path, old_lines_pair, new_lines_pair, out_path, ts))

        saved, cancelled, failures = _run_parallel_with_progress(
            parent, items, _save_one_diff_html, "HTML")

        if cancelled:
            status_bar.showMessage(
                f"⚠  HTML 저장 취소됨 ({saved}/{len(checked_pairs)}개 저장): {out_dir}")
        elif failures:
            status_bar.showMessage(
                f"⚠  HTML 저장 일부 실패 ({saved}/{len(checked_pairs)}개 성공): {out_dir}")
        else:
            status_bar.showMessage(
                f"✅  HTML 저장 완료 ({saved}개 파일, 개별 저장): {out_dir}")
        return

    # ── 기존 동작: 현재 선택된 단일 파일 or 전체 병합 ──
    cur = result_panel.get_current_file_lines()
    if cur:
        rel_path, old_lines, new_lines = cur
        fname = rel_path.replace("\\", "/").split("/")[-1]
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Diff_{fname}_{ts}.html"
        page_title = f"🔀 코드 변경점 — {fname}"
    else:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Full_Diff_{ts}.html"
        page_title = "🔀 코드 변경점 전체 보기"

    path, _ = QFileDialog.getSaveFileName(
        parent, "코드 변경점 HTML 저장", default_name, "HTML (*.html)")
    if not path:
        return

    dlg = QProgressDialog("HTML 저장 중...", None, 0, 0, parent)
    dlg.setWindowModality(Qt.WindowModality.WindowModal)
    dlg.setCancelButton(None)
    dlg.setMinimumDuration(0)
    dlg.show()
    QApplication.processEvents()
    try:
        rows_html_joined = build_html_diff_rows(old_lines, new_lines)
        body_block = f"""
<div class="wrap">
<table>
<colgroup><col><col><col><col></colgroup>
<thead>
  <tr>
    <th>행#</th><th>수정 전 코드</th>
    <th>행#</th><th>수정 후 코드</th>
  </tr>
</thead>
<tbody>
{rows_html_joined}
</tbody>
</table>
</div>"""
        html = html_shell(page_title, ts, body_block)
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
    finally:
        dlg.close()
    status_bar.showMessage(f"✅  HTML 저장 완료: {path}")
